"""
Exportación a Excel del panel del profesor.
"""

from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1A1A2E")
HEADER_FONT = Font(bold=True, color="FFFFFF")

# Formato por nombre de columna
MONEY = '"$"#,##0'
PRICE = '"$"#,##0.00'
PCT = '0.00"%"'
DEC2 = "0.00"
QTY = "#,##0.0000"
FORMATS = {
    "Invertido": MONEY, "Efectivo": MONEY, "Valor Total": MONEY, "Valor (COP)": MONEY, "Monto": MONEY,
    "Invertido (COP)": MONEY, "Efectivo (COP)": MONEY, "Valor total (COP)": MONEY,
    "Precio": PRICE, "Cantidad": QTY,
    "Return (%)": PCT, "% Efectivo": PCT, "Peso (%)": PCT, "Volatilidad anual (%)": PCT,
    "Alpha anual (%)": PCT, "Max drawdown (%)": PCT, "Tracking error (%)": PCT,
    "Sharpe": DEC2, "Beta": DEC2, "Information ratio": DEC2,
}


def _write(writer, df: pd.DataFrame, sheet: str, empty_msg: str):
    if df is None or df.empty:
        df = pd.DataFrame({"Nota": [empty_msg]})
    df = df.copy()
    for col in df.columns:
        if pd.api.types.is_datetime64_any_dtype(df[col]):
            df[col] = df[col].dt.tz_localize(None) if getattr(df[col].dt, "tz", None) else df[col]
    df.to_excel(writer, sheet_name=sheet, index=False)
    ws = writer.sheets[sheet]
    # La exportación nunca escribe fórmulas: cualquier texto que Excel fuera a
    # interpretar como fórmula (p. ej. un nickname que empieza por "=") queda como texto.
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                cell.data_type = "s"
    ws.freeze_panes = "A2"
    for idx, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        fmt = FORMATS.get(str(col))
        if pd.api.types.is_datetime64_any_dtype(df[col]):
            fmt = "dd/mm/yyyy"
        if fmt:
            for row in range(2, len(df) + 2):
                ws.cell(row=row, column=idx).number_format = fmt
        longest = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(10, longest + 2), 45)
    ws.row_dimensions[1].height = 30


def build_excel(leaderboard: pd.DataFrame, metrics: pd.DataFrame, positions: pd.DataFrame,
                trades: pd.DataFrame, snapshots: pd.DataFrame, groups: dict) -> bytes:
    """Devuelve el archivo .xlsx (bytes) con una hoja por tema."""
    snaps = pd.DataFrame()
    if snapshots is not None and not snapshots.empty:
        nick = {k: v.get("nickname", "") for k, v in groups.items()}
        snaps = snapshots.copy()
        snaps.insert(2, "Nickname", snaps["group_number"].map(nick))
        snaps = snaps.rename(columns={
            "date": "Fecha", "group_number": "Grupo", "invested": "Invertido (COP)",
            "cash": "Efectivo (COP)", "total_value": "Valor total (COP)",
            "return_pct": "Return (%)", "cash_pct": "% Efectivo",
        }).sort_values(["Fecha", "Grupo"])

    lb = leaderboard
    if lb is not None and not lb.empty and "Pos" not in lb.columns and "index" in lb.columns:
        lb = lb.rename(columns={"index": "Pos"})

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        _write(writer, lb, "Leaderboard", "Sin grupos registrados")
        _write(writer, metrics, "Métricas de riesgo", "Sin datos suficientes")
        _write(writer, positions, "Posiciones", "Sin posiciones")
        _write(writer, trades, "Operaciones", "Sin operaciones")
        _write(writer, snaps, "Fotos diarias", "Aún no hay fotos del cierre (las guarda el reporte nocturno)")
    return buf.getvalue()
